"""Reporte del análisis en PDF y Word: malla, esfuerzos de Von Mises y deformada."""

import shutil
from pathlib import Path

import numpy as np
import pytest

from ifc_structural.core.cargas import ApoyoSuperficie
from ifc_structural.core.casos_carga import CasoCarga, Combinacion, TipoCaso
from ifc_structural.reporte import DatosReporte, armar_contenido, exportar_reporte
from ifc_structural.solido.analisis_solido import ejecutar_analisis_solido, piezas_con_mecanismo

from .test_solidos_libres import _carga, caja, cara_hacia, proyecto_con

requiere_ccx = pytest.mark.skipif(shutil.which("ccx") is None, reason="ccx no está instalado en este sistema")


@pytest.fixture(scope="module")
def analisis(tmp_path_factory):
    s = caja("S1", 0, 0, 0, 1.0, 0.1, 0.1)
    p = proyecto_con(s)
    apoyo = ApoyoSuperficie(id="AP1", elemento_tipo="SOLIDO", elemento_id="S1", cara=cara_hacia(s, (-1, 0, 0)),
                            ux=True, uy=True, uz=True)
    casos = {"D": CasoCarga("D", "Muerta", TipoCaso.D), "L": CasoCarga("L", "Viva", TipoCaso.L)}
    combinaciones = {"C1": Combinacion("C1", "1.2D+1.6L", {"D": 1.2, "L": 1.6})}
    carga = _carga(s, (1, 0, 0), "FUERZA", 10.0, (0, 0, -1))
    carga.caso_id = "L"
    r = ejecutar_analisis_solido(p, casos, combinaciones, [apoyo], {}, {}, {}, tmp_path_factory.mktemp("ccx"), tamano=0.05,
                                 cargas_superficie_por_caso={"L": [carga]})
    return DatosReporte(p, r, casos, combinaciones, [apoyo], {"D": 0, "L": 1})


@requiere_ccx
def test_el_contenido_resume_malla_von_mises_y_deformada(analisis):
    imagenes = {"malla": "m.png", "vm:L": "v.png", "u:L": "u.png"}
    bloques = armar_contenido(analisis, imagenes)
    textos = " ".join(b.texto for b in bloques)
    assert "Malla" in textos and "C3D10" in textos
    resumen = next(b for b in bloques if b.tipo == "tabla" and b.filas[0][0] == "Resultado")
    assert [f[0] for f in resumen.filas[1:]] == ["Caso D", "Caso L", "Combinación C1"]
    # σvm máximo del voladizo: cerca del empotramiento, del orden de M·c/I = F·L·(h/2)/(b h³/12)
    vm_l = float(resumen.filas[2][1].replace(" ", ""))
    assert 0.3 * 60.0 < vm_l < 3 * 60.0
    assert {b.ruta for b in bloques if b.tipo == "imagen"} == {"m.png", "v.png", "u.png"}


@requiere_ccx
@pytest.mark.parametrize("extension", ["pdf", "docx"])
def test_exporta_el_reporte_con_las_imagenes(analisis, tmp_path: Path, extension):
    ruta = exportar_reporte(tmp_path / f"reporte.{extension}", analisis)
    assert ruta.exists() and ruta.stat().st_size > 50_000
    if extension == "pdf":
        contenido = ruta.read_bytes()
        assert contenido.startswith(b"%PDF") and contenido.count(b"/Subtype /Image") >= 7  # malla + 2 por resultado
    else:
        from docx import Document

        doc = Document(str(ruta))
        titulos = [p.text for p in doc.paragraphs if p.style.name.startswith(("Heading", "Title"))]
        assert titulos[0] == "Reporte de análisis estructural"
        assert any("Malla" in t for t in titulos) and any("Combinación C1" in t for t in titulos)
        assert len(doc.inline_shapes) == 7


def test_solo_pdf_o_word(analisis, tmp_path: Path):
    with pytest.raises(ValueError):
        exportar_reporte(tmp_path / "reporte.txt", analisis)


def test_detecta_piezas_que_forman_un_mecanismo():
    class Malla:
        coordenadas = np.array([[0, 0, 0], [1, 0, 0], [0, 1, 0], [5, 0, 0], [6, 0, 0], [5, 1, 0]], dtype=float)
        conectividad = np.array([[0, 1, 2, 2, 2, 2, 2, 2], [3, 4, 5, 5, 5, 5, 5, 5]])

    u = np.zeros((6, 3))
    assert piezas_con_mecanismo(Malla, {"D": u}, ["A", "B"]) == []
    u[3:, 2] = 1e6  # la pieza B se va lejos: la matriz era casi singular
    assert piezas_con_mecanismo(Malla, {"D": u}, ["A", "B"]) == ["B"]


@requiere_ccx
def test_reporte_de_un_analisis_de_barras(tmp_path: Path):
    from ifc_structural.core.cargas import Apoyo, CargaArea
    from ifc_structural.ifc.importador import importar_ifc
    from ifc_structural.solver.analisis import ejecutar_analisis

    from .ayuda_ifc import crear_ifc_portico
    from .test_analisis import CASOS, COMBINACIONES

    crear_ifc_portico(str(tmp_path / "portico.ifc"))
    p = importar_ifc(str(tmp_path / "portico.ifc"))
    apoyos = [Apoyo(id=f"AP{i}", nodo_id=c.nodo_inicio) for i, c in enumerate(p.columnas.values(), start=1)]
    losa = next(iter(p.losas.values()))
    r = ejecutar_analisis(p, CASOS, COMBINACIONES, apoyos, {"L": [CargaArea(id="CA1", caso_id="L", presion=5.0, losa_id=losa.id)]},
                          {}, {}, tmp_path / "trabajo")
    datos = DatosReporte(p, r, CASOS, COMBINACIONES, apoyos, {"L": 1})
    bloques = armar_contenido(datos, {})
    assert any("modelo de barras" in b.texto for b in bloques)
    ruta = exportar_reporte(tmp_path / "portico.docx", datos)
    from docx import Document

    assert len(Document(str(ruta)).inline_shapes) == 1 + 2 * (len(r.tensiones_por_caso) + len(r.tensiones_por_combinacion))
