"""Reporte del análisis en PDF o Word: datos del modelo, malla, esfuerzos de Von Mises y deformada.

El contenido se arma una sola vez como una lista de bloques (títulos, párrafos, tablas e imágenes) y se escribe con
``reportlab`` (PDF) o con ``python-docx`` (Word), así los dos formatos dicen exactamente lo mismo. Las imágenes se
dibujan con PyVista fuera de pantalla, independientes del visor (no cambian la vista del usuario).
"""

from __future__ import annotations

import datetime
import math
import os
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Dict, List, Optional, Sequence, Tuple

import numpy as np

from .combinaciones import esfuerzos
from .core.casos_carga import CasoCarga, Combinacion
from .core.modelo import Proyecto
from .solver.analisis import ResultadoAnalisis

#: Tamaño de las imágenes (píxeles) y ancho con que se colocan en la página (cm).
TAMANO_IMAGEN = (1600, 1000)
ANCHO_IMAGEN_CM = 16.5
#: Como mucho se dibujan tantos casos y combinaciones (los demás solo van en la tabla de resumen).
MAXIMO_RESULTADOS_CON_IMAGEN = 12


# ---------------------------------------------------------------------------
# Contenido
# ---------------------------------------------------------------------------


@dataclass
class Bloque:
    tipo: str  # "titulo", "subtitulo", "parrafo", "tabla", "imagen", "salto"
    texto: str = ""
    filas: List[List[str]] = field(default_factory=list)  # tabla: la primera fila es la cabecera
    ruta: str = ""  # imagen
    nivel: int = 1


@dataclass
class DatosReporte:
    proyecto: Proyecto
    resultado: ResultadoAnalisis
    casos: Dict[str, CasoCarga]
    combinaciones: Dict[str, Combinacion]
    apoyos: Sequence = ()
    cargas_por_caso: Dict[str, int] = field(default_factory=dict)  # caso -> número de cargas asignadas
    escala_deformada: Optional[float] = None  # None = automática
    colormap: str = "turbo"
    autor: str = ""


@dataclass
class Resumen:
    id: str
    nombre: str
    tipo: str  # "Caso" o "Combinación"
    vm_max: float  # MPa
    punto_vm: np.ndarray
    elemento_vm: str
    u_max: float  # mm
    punto_u: np.ndarray
    u_componentes: Tuple[float, float, float]  # máximos de |ux|, |uy|, |uz| en mm


def _fmt(x: float, d: int = 3) -> str:
    return f"{x:,.{d}f}".replace(",", " ")


def _punto(p: np.ndarray) -> str:
    return f"({p[0]:.3f}; {p[1]:.3f}; {p[2]:.3f})"


def _resultados(datos: DatosReporte) -> List[Tuple[str, str, str, np.ndarray, Optional[np.ndarray]]]:
    """(id, nombre, tipo, tensiones, desplazamientos) de cada caso y combinación analizados."""

    r = datos.resultado
    lista = []
    for id_, t in r.tensiones_por_caso.items():
        caso = datos.casos.get(id_)
        lista.append((id_, caso.nombre if caso else id_, "Caso", t, r.desplazamientos_por_caso.get(id_)))
    for id_, t in r.tensiones_por_combinacion.items():
        comb = datos.combinaciones.get(id_)
        lista.append((id_, comb.nombre if comb else id_, "Combinación", t, r.desplazamientos_por_combinacion.get(id_)))
    return lista


def _resumenes(datos: DatosReporte) -> List[Resumen]:
    r = datos.resultado
    malla = r.malla_tensiones
    coordenadas = np.asarray(malla.coordenadas, dtype=float)
    # una celda por nodo, para decir a qué elemento pertenece el máximo
    celda_de_nodo = np.full(len(coordenadas), -1, dtype=np.int64)
    conectividad = np.asarray(malla.conectividad, dtype=np.int64)
    celda_de_nodo[conectividad.ravel()] = np.repeat(np.arange(len(conectividad)), conectividad.shape[1])
    barras = r.barra_de_cada_celda()
    resumenes = []
    for id_, nombre, tipo, t, u in _resultados(datos):
        vm = esfuerzos.von_mises(t) * esfuerzos.FACTOR_KPA_A_MPA
        i = int(np.nanargmax(vm))
        celda = int(celda_de_nodo[i])
        if u is not None and len(u):
            norma = np.linalg.norm(u, axis=1)
            j = int(np.nanargmax(norma))
            u_max, punto_u = float(norma[j]) * 1000.0, coordenadas[j]
            componentes = tuple(float(np.nanmax(np.abs(u[:, k]))) * 1000.0 for k in range(3))
        else:
            u_max, punto_u, componentes = float("nan"), coordenadas[0] * np.nan, (float("nan"),) * 3
        resumenes.append(Resumen(id_, nombre, tipo, float(vm[i]), coordenadas[i], barras[celda] if celda >= 0 else "—",
                                 u_max, punto_u, componentes))
    return resumenes


def escala_deformada_automatica(resultado: ResultadoAnalisis) -> float:
    """Ampliación con la que el mayor desplazamiento mide ~8 % del tamaño del modelo (la misma que el visor)."""

    todos = list(resultado.desplazamientos_por_caso.values()) + list(resultado.desplazamientos_por_combinacion.values())
    if not todos:
        return 1.0
    u_max = max(float(np.linalg.norm(u, axis=1).max()) for u in todos)
    c = resultado.malla_tensiones.coordenadas
    tamano = float(np.linalg.norm(c.max(axis=0) - c.min(axis=0)))
    if u_max < 1e-12 or tamano < 1e-9:
        return 1.0
    bruta = 0.08 * tamano / u_max
    cifras = 10 ** (math.floor(math.log10(bruta)) - 1)
    return float(min(1e7, max(1.0, round(bruta / cifras) * cifras)))


def _linea_de_malla(resultado: ResultadoAnalisis) -> str:
    return next((a for a in resultado.advertencias if a.startswith("Malla:")), "")


def armar_contenido(datos: DatosReporte, imagenes: Dict[str, str]) -> List[Bloque]:
    """Bloques del reporte; ``imagenes`` da la ruta de cada imagen ("malla", "vm:<id>", "u:<id>")."""

    p, r = datos.proyecto, datos.resultado
    bloques: List[Bloque] = []
    fecha = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
    bloques.append(Bloque("titulo", "Reporte de análisis estructural"))
    bloques.append(Bloque("parrafo", f"Proyecto: {p.nombre or 'sin nombre'}" + (f" ({Path(p.ruta_ifc).name})" if getattr(p, "ruta_ifc", None) else "")))
    bloques.append(Bloque("parrafo", f"Fecha: {fecha}" + (f" — Autor: {datos.autor}" if datos.autor else "")))
    bloques.append(Bloque("parrafo", "Tipo de análisis: " + ("modelo sólido 3D por elementos finitos (CalculiX)" if r.es_solido
                                                           else "modelo de barras (CalculiX)") + ", elástico lineal."))

    # 1. Modelo
    bloques.append(Bloque("subtitulo", "1. Modelo"))
    filas = [["Elemento", "Cantidad"]]
    for nombre, n in (("Vigas", len(p.vigas)), ("Columnas", len(p.columnas)), ("Losas", len(p.losas)), ("Muros", len(p.muros)),
                      ("Sólidos de forma libre", len(p.solidos))):
        if n:
            filas.append([nombre, str(n)])
    bloques.append(Bloque("tabla", filas=filas))
    usados = {getattr(e, "material", None) for c in (p.vigas, p.columnas, p.losas, p.muros, p.solidos) for e in c.values()}
    filas = [["Material", "E [MPa]", "ν", "Peso específico [kN/m³]"]]
    for nombre, m in p.materiales.items():
        if nombre in usados:
            filas.append([nombre, _fmt(m.modulo_elasticidad / 1000.0, 0), f"{m.modulo_poisson:.2f}", f"{m.densidad:.2f}"])
    if len(filas) > 1:
        bloques.append(Bloque("parrafo", "Materiales:"))
        bloques.append(Bloque("tabla", filas=filas))

    # 2. Apoyos y cargas
    bloques.append(Bloque("subtitulo", "2. Apoyos, cargas y combinaciones"))
    bloques.append(Bloque("parrafo", f"Apoyos asignados: {len(datos.apoyos)}."))
    if datos.apoyos:
        nombres_gdl = {1: "Ux", 2: "Uy", 3: "Uz", 4: "Rx", 5: "Ry", 6: "Rz"}
        filas = [["Apoyo", "Sobre", "Restringe"]]
        for a in datos.apoyos:
            if getattr(a, "nodo_id", None):
                sobre = f"nodo {a.nodo_id}"
            else:
                sobre = f"{str(getattr(a, 'elemento_tipo', '')).lower()} {getattr(a, 'elemento_id', '')}".strip()
                if getattr(a, "cara", None) is not None:
                    sobre += f", superficie {a.cara}"
            grados = a.grados_restringidos() if hasattr(a, "grados_restringidos") else []
            filas.append([getattr(a, "nombre", "") or getattr(a, "id", ""), sobre, " ".join(nombres_gdl.get(g, str(g)) for g in grados)])
        bloques.append(Bloque("tabla", filas=filas))
    filas = [["Caso", "Nombre", "Tipo", "Cargas"]]
    for c in datos.casos.values():
        tipo = getattr(c.tipo, "value", str(c.tipo)) + (" + peso propio" if getattr(c, "incluye_peso_propio", False) else "")
        filas.append([c.id, c.nombre, str(tipo), str(datos.cargas_por_caso.get(c.id, 0))])
    bloques.append(Bloque("tabla", filas=filas))
    if datos.combinaciones:
        filas = [["Combinación", "Nombre", "Factores"]]
        for c in datos.combinaciones.values():
            factores = getattr(c, "factores", {}) or {}
            filas.append([c.id, c.nombre, " + ".join(f"{f:g}·{k}" for k, f in factores.items())])
        bloques.append(Bloque("tabla", filas=filas))

    # 3. Malla
    malla = r.malla_tensiones
    bloques.append(Bloque("subtitulo", "3. Malla de elementos finitos"))
    texto = _linea_de_malla(r) or f"{len(malla.conectividad):,} elementos, {len(malla.coordenadas):,} nodos.".replace(",", " ")
    bloques.append(Bloque("parrafo", texto))
    if r.es_solido:
        bloques.append(Bloque("parrafo", "Elementos: tetraedros cuadráticos C3D10 en las piezas macizas y láminas S6 en las "
                                         "piezas que son láminas (espesor ≤ 1/20 de su tamaño en planta). Las láminas se dibujan "
                                         "con su espesor."))
    if "malla" in imagenes:
        bloques.append(Bloque("imagen", "Malla del modelo", ruta=imagenes["malla"]))

    # 4. Resultados
    resumenes = _resumenes(datos)
    escala = datos.escala_deformada or escala_deformada_automatica(r)
    bloques.append(Bloque("subtitulo", "4. Resumen de resultados"))
    filas = [["Resultado", "σvm máx. [MPa]", "Ubicación (x; y; z) [m]", "Elemento", "Desplaz. máx. [mm]"]]
    for s in resumenes:
        filas.append([f"{s.tipo} {s.id}", _fmt(s.vm_max), _punto(s.punto_vm), s.elemento_vm,
                      "—" if math.isnan(s.u_max) else _fmt(s.u_max)])
    bloques.append(Bloque("tabla", filas=filas))
    peor = max(resumenes, key=lambda s: s.vm_max, default=None)
    if peor is not None:
        bloques.append(Bloque("parrafo", f"Esfuerzo de Von Mises más desfavorable: {_fmt(peor.vm_max)} MPa ({peor.tipo.lower()} "
                                         f"{peor.id}, elemento {peor.elemento_vm}). Los valores junto a apoyos, cargas puntuales y "
                                         "esquinas interiores son locales y dependen del tamaño de elemento."))

    for k, s in enumerate(resumenes):
        bloques.append(Bloque("salto"))
        bloques.append(Bloque("subtitulo", f"5.{k + 1} {s.tipo} {s.id} — {s.nombre}", nivel=2))
        bloques.append(Bloque("parrafo", f"Von Mises máximo: {_fmt(s.vm_max)} MPa en {_punto(s.punto_vm)} m (elemento "
                                         f"{s.elemento_vm})."))
        if f"vm:{s.id}" in imagenes:
            bloques.append(Bloque("imagen", f"Esfuerzos de Von Mises [MPa] — {s.tipo.lower()} {s.id}", ruta=imagenes[f"vm:{s.id}"]))
        if not math.isnan(s.u_max):
            ux, uy, uz = s.u_componentes
            bloques.append(Bloque("parrafo", f"Desplazamiento máximo: {_fmt(s.u_max)} mm en {_punto(s.punto_u)} m "
                                             f"(|ux| máx. {_fmt(ux)}, |uy| máx. {_fmt(uy)}, |uz| máx. {_fmt(uz)} mm)."))
            if f"u:{s.id}" in imagenes:
                bloques.append(Bloque("imagen", f"Deformada ampliada ×{escala:g}, coloreada por el desplazamiento [mm] — "
                                                f"{s.tipo.lower()} {s.id}", ruta=imagenes[f"u:{s.id}"]))
    if len(resumenes) > MAXIMO_RESULTADOS_CON_IMAGEN:
        bloques.append(Bloque("parrafo", f"Solo se dibujan los primeros {MAXIMO_RESULTADOS_CON_IMAGEN} resultados; el resto figura "
                                         "en la tabla de resumen."))

    otras = [a for a in r.advertencias if not a.startswith("Malla:")]
    if otras:
        bloques.append(Bloque("salto"))
        bloques.append(Bloque("subtitulo", "6. Advertencias del análisis"))
        for a in otras:
            bloques.append(Bloque("parrafo", f"• {a}"))
    return bloques


# ---------------------------------------------------------------------------
# Imágenes
# ---------------------------------------------------------------------------


def generar_imagenes(datos: DatosReporte, carpeta: Path,
                     al_avanzar: Optional[Callable[[str], None]] = None) -> Dict[str, str]:
    """Imágenes PNG de la malla y, por caso y combinación, de los esfuerzos de Von Mises y de la deformada."""

    import warnings

    import pyvista as pv

    r = datos.resultado
    malla = r.malla_tensiones
    coordenadas = np.asarray(malla.coordenadas, dtype=float)
    celdas, tipos = malla.celdas_vtk()
    grid = pv.UnstructuredGrid(celdas, tipos, coordenadas.copy())
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        superficie = grid.extract_surface(pass_pointid=True)
    ids = np.asarray(superficie.point_data["vtkOriginalPointIds"], dtype=np.int64)
    contorno = superficie.extract_feature_edges(feature_angle=35, boundary_edges=True, manifold_edges=False,
                                                non_manifold_edges=False)
    escala = datos.escala_deformada or escala_deformada_automatica(r)
    imagenes: Dict[str, str] = {}

    def avisar(texto: str) -> None:
        if al_avanzar is not None:
            al_avanzar(texto)

    def foto(nombre: str, dibujar) -> str:
        pl = pv.Plotter(off_screen=True, window_size=TAMANO_IMAGEN)
        try:
            pl.set_background("white")
            dibujar(pl)
            pl.view_isometric()
            pl.reset_camera()
            pl.camera.zoom(1.15)
            ruta = str(carpeta / f"{nombre}.png")
            pl.screenshot(ruta)
        finally:
            pl.close()
        return ruta

    barra = dict(title_font_size=18, label_font_size=16, color="black", fmt="%.3g", vertical=True,
                 position_x=0.88, position_y=0.15, height=0.7, width=0.06)

    avisar("Reporte: dibujando la malla...")
    imagenes["malla"] = foto("malla", lambda pl: pl.add_mesh(superficie, color="#c9d6e3", show_edges=True,
                                                             edge_color="#1f2a36", line_width=0.4))
    for k, (id_, _, tipo, t, u) in enumerate(_resultados(datos)[:MAXIMO_RESULTADOS_CON_IMAGEN]):
        avisar(f"Reporte: dibujando {tipo.lower()} {id_}...")
        vm = esfuerzos.von_mises(t) * esfuerzos.FACTOR_KPA_A_MPA
        sup = superficie.copy()
        sup.point_data["vm"] = vm[ids]

        def dibujar_vm(pl, sup=sup):
            pl.add_mesh(sup, scalars="vm", cmap=datos.colormap, scalar_bar_args=dict(title="Von Mises [MPa]", **barra))
            pl.add_mesh(contorno, color="#333333", line_width=1.0)

        imagenes[f"vm:{id_}"] = foto(f"vm_{k}", dibujar_vm)
        if u is None or not len(u):
            continue
        deformada = superficie.copy()
        deformada.points = coordenadas[ids] + escala * np.asarray(u, dtype=float)[ids]
        deformada.point_data["u"] = np.linalg.norm(u, axis=1)[ids] * 1000.0

        def dibujar_u(pl, deformada=deformada):
            pl.add_mesh(contorno, color="#8a96a3", line_width=1.0, opacity=0.8)
            pl.add_mesh(deformada, scalars="u", cmap=datos.colormap,
                        scalar_bar_args=dict(title="Desplazamiento [mm]", **barra))

        imagenes[f"u:{id_}"] = foto(f"u_{k}", dibujar_u)
    return imagenes


# ---------------------------------------------------------------------------
# Escritura
# ---------------------------------------------------------------------------


def _fuente_unicode() -> Optional[str]:
    """Ruta de una fuente TrueType con σ, ν, ≤ y ³ (DejaVu Sans, que trae matplotlib)."""

    try:
        import matplotlib

        ruta = Path(matplotlib.get_data_path()) / "fonts" / "ttf" / "DejaVuSans.ttf"
        return str(ruta) if ruta.exists() else None
    except Exception:  # noqa: BLE001
        return None


def escribir_pdf(bloques: Sequence[Bloque], ruta: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    fuente, negrita = "Helvetica", "Helvetica-Bold"
    ttf = _fuente_unicode()
    if ttf:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        try:
            pdfmetrics.registerFont(TTFont("DejaVu", ttf))
            bold = Path(ttf).with_name("DejaVuSans-Bold.ttf")
            pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(bold) if bold.exists() else ttf))
            fuente, negrita = "DejaVu", "DejaVu-Bold"
        except Exception:  # noqa: BLE001 - se queda con Helvetica
            pass
    base = getSampleStyleSheet()
    estilos = {
        "titulo": ParagraphStyle("t", parent=base["Title"], fontName=negrita, fontSize=20, spaceAfter=12),
        "subtitulo": ParagraphStyle("s", parent=base["Heading2"], fontName=negrita, fontSize=14, spaceBefore=10, spaceAfter=6,
                                    textColor=colors.HexColor("#1f3a5f")),
        "parrafo": ParagraphStyle("p", parent=base["BodyText"], fontName=fuente, fontSize=10, leading=13),
        "pie": ParagraphStyle("pie", parent=base["BodyText"], fontName=fuente, fontSize=9, alignment=TA_CENTER,
                              textColor=colors.HexColor("#555555"), spaceAfter=8),
        "celda": ParagraphStyle("c", parent=base["BodyText"], fontName=fuente, fontSize=8.5, leading=10),
        "cabecera": ParagraphStyle("h", parent=base["BodyText"], fontName=negrita, fontSize=8.5, leading=10,
                                   textColor=colors.white),
    }
    historia = []
    ancho = A4[0] - 3.6 * cm
    for b in bloques:
        if b.tipo == "titulo":
            historia.append(Paragraph(escape(b.texto), estilos["titulo"]))
        elif b.tipo == "subtitulo":
            historia.append(Paragraph(escape(b.texto), estilos["subtitulo"]))
        elif b.tipo == "parrafo":
            historia.append(Paragraph(escape(b.texto), estilos["parrafo"]))
        elif b.tipo == "salto":
            historia.append(PageBreak())
        elif b.tipo == "tabla" and b.filas:
            datos = [[Paragraph(escape(c), estilos["cabecera"] if i == 0 else estilos["celda"]) for c in fila]
                     for i, fila in enumerate(b.filas)]
            tabla = Table(datos, repeatRows=1, hAlign="LEFT", colWidths=[ancho / len(b.filas[0])] * len(b.filas[0]))
            tabla.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f3a5f")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef2f7")]),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa8b8")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]))
            historia += [tabla, Spacer(1, 8)]
        elif b.tipo == "imagen" and b.ruta:
            w, h = ImageReader(b.ruta).getSize()
            ancho_img = min(ancho, ANCHO_IMAGEN_CM * cm)
            historia += [Image(b.ruta, width=ancho_img, height=ancho_img * h / w), Paragraph(escape(b.texto), estilos["pie"])]

    def pie_de_pagina(canvas, doc):
        canvas.saveState()
        canvas.setFont(fuente, 8)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawString(1.8 * cm, 1.2 * cm, "Analizador estructural IFC — reporte de análisis")
        canvas.drawRightString(A4[0] - 1.8 * cm, 1.2 * cm, f"Página {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(str(ruta), pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm, topMargin=1.8 * cm,
                            bottomMargin=2.0 * cm, title="Reporte de análisis estructural")
    doc.build(historia, onFirstPage=pie_de_pagina, onLaterPages=pie_de_pagina)


def escribir_word(bloques: Sequence[Bloque], ruta: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    for seccion in doc.sections:
        seccion.left_margin = seccion.right_margin = Cm(1.8)
        seccion.top_margin = seccion.bottom_margin = Cm(1.8)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    def sombrear(celda, color: str) -> None:
        propiedades = celda._tc.get_or_add_tcPr()
        sombra = OxmlElement("w:shd")
        sombra.set(qn("w:val"), "clear")
        sombra.set(qn("w:color"), "auto")
        sombra.set(qn("w:fill"), color)
        propiedades.append(sombra)

    for b in bloques:
        if b.tipo == "titulo":
            doc.add_heading(b.texto, level=0)
        elif b.tipo == "subtitulo":
            doc.add_heading(b.texto, level=1 if b.nivel == 1 else 2)
        elif b.tipo == "parrafo":
            doc.add_paragraph(b.texto)
        elif b.tipo == "salto":
            doc.add_page_break()
        elif b.tipo == "tabla" and b.filas:
            tabla = doc.add_table(rows=len(b.filas), cols=len(b.filas[0]))
            tabla.style = "Table Grid"
            for i, fila in enumerate(b.filas):
                for j, texto in enumerate(fila):
                    celda = tabla.cell(i, j)
                    celda.text = texto
                    for parrafo in celda.paragraphs:
                        for run in parrafo.runs:
                            run.font.size = Pt(9)
                            if i == 0:
                                run.bold = True
                                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    if i == 0:
                        sombrear(celda, "1F3A5F")
                    elif i % 2 == 0:
                        sombrear(celda, "EEF2F7")
            doc.add_paragraph()
        elif b.tipo == "imagen" and b.ruta:
            doc.add_picture(b.ruta, width=Cm(ANCHO_IMAGEN_CM))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            pie = doc.add_paragraph(b.texto)
            pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in pie.runs:
                run.italic = True
                run.font.size = Pt(9)
    doc.save(str(ruta))


def exportar_reporte(ruta, datos: DatosReporte, al_avanzar: Optional[Callable[[str], None]] = None) -> Path:
    """Escribe el reporte en ``ruta``: PDF si termina en .pdf, Word si termina en .docx."""

    ruta = Path(ruta)
    sufijo = ruta.suffix.lower()
    if sufijo not in (".pdf", ".docx"):
        raise ValueError("El reporte se exporta a PDF (.pdf) o a Word (.docx).")
    with tempfile.TemporaryDirectory(prefix="reporte_", ignore_cleanup_errors=True) as carpeta:
        imagenes = generar_imagenes(datos, Path(carpeta), al_avanzar)
        bloques = armar_contenido(datos, imagenes)
        if al_avanzar is not None:
            al_avanzar("Reporte: escribiendo el documento...")
        if sufijo == ".pdf":
            escribir_pdf(bloques, ruta)
        else:
            escribir_word(bloques, ruta)
    return ruta
